#!/usr/bin/env python3
"""Turns the programmes workbook into the file the app reads.

    python3 scripts/import_programmes.py ["programmes/Kinetix Programmes.xlsx"]

Reads the Programmes, Run weeks, Run sessions and Rules tabs, works out each workout's
steps from the way it's written ("6 x 800 m @ I, 90 s jog between"), and writes
Packages/TrainingEngine/Sources/TrainingEngine/Resources/programmes.json.

Anything it can't understand is listed with its tab and row, and nothing is written, so
the app never gets half a programme. Distances written in a workout ("13 km E, final
3 km @ STEADY") set the shape of the run; the app sizes it from the runner's volume.
"""
import json
import re
import sys
import warnings
from pathlib import Path

warnings.filterwarnings("ignore")  # openpyxl: Excel's extended drop-downs aren't needed here
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_WORKBOOK = ROOT / "programmes" / "Kinetix Programmes.xlsx"
OUTPUT = ROOT / "Packages/TrainingEngine/Sources/TrainingEngine/Resources/programmes.json"

GOALS = {
    "first 5k": "first5k", "faster 5k": "faster5k", "faster 10k": "faster10k", "half marathon": "halfMarathon", "marathon": "marathon",
    "build strength": "buildStrength", "build muscle": "buildMuscle", "balanced hybrid": "balancedHybrid", "hyrox": "hyrox",
}
# Workbook zone codes → app zones.
ZONES = {
    "e": "easy", "easy": "easy", "steady": "steady", "m": "marathon", "mp": "marathon", "hm": "halfMarathon",
    "hmp": "halfMarathon", "t": "threshold", "10k": "tenK", "5k": "fiveK", "i": "interval", "r": "repetition",
    # Coach's words: tempo is threshold pace; 3K is run at interval pace (the app has no 3K zone).
    "tempo": "threshold", "3k": "interval",
}
# Longer codes first, so "EASY" isn't read as "E" and "TEMPO" isn't read as "T".
ZONE_PATTERN = r"(steady|easy|tempo|hmp|hm|mp|m|10k|5k|3k|t|i|r|e)"
# Session type column → how the app schedules it, and the name shown.
SESSION_TYPES = {
    "easy": ("easy", "Easy run"), "e": ("easy", "Easy run"), "recovery": ("recovery", "Recovery run"),
    "long": ("long", "Long run"), "tempo": ("tempo", "Tempo run"), "progression": ("tempo", "Progression run"),
    "intervals": ("intervals", "Intervals"), "fartlek": ("intervals", "Fartlek"), "hills": ("intervals", "Hills"),
    "shakeout": ("easy", "Shakeout"), "race": ("race", "Race day"),
    # Hyrox: runs with station work between the reps (e.g. '6 x 1 km @ HM, 2 min station between').
    "compromised": ("intervals", "Compromised run"),
}
SPLITS = {
    "full body ×1": "fullBody1", "full body ×2": "fullBody2", "full body ×3": "fullBody3", "upper / lower": "upperLower",
    "push / pull / legs": "ppl", "push / pull / legs ×2": "ppl6", "upper / lower / full": "upperLowerFull",
    "ppl + upper / lower": "pplUpperLower", "ppl + upper + shoulders & arms": "pplUpperArms",
    "push / pull / legs / upper": "pplUpper",
    "hyrox: stations ×2": "hyroxStations2", "hyrox: strength + stations": "hyroxHybrid3",
}
PATTERNS = {
    "any squat": "squat", "any hinge (deadlift)": "hinge", "any hinge accessory (rdl, bridge)": "hingeAccessory",
    "any lunge / split squat": "lunge", "any calf raise": "calf", "any horizontal push (bench, press-up)": "horizontalPush",
    "any vertical push (overhead press)": "verticalPush", "any horizontal pull (row)": "horizontalPull",
    "any vertical pull (pull-up, pulldown)": "verticalPull", "any arms": "arms", "any core": "core",
    "any anti-rotation core": "antiRotation", "any hyrox station": "station",
}
EXERCISES = ROOT / "Packages/TrainingEngine/Sources/TrainingEngine/Resources/exercises.json"
PHASES = {"base": "base", "build": "build", "peak": "peak", "taper": "taper", "race week": "taper", "deload": None}
EMPHASIS = {"normal": "normal", "heavy": "heavy", "light": "light", "none": "none"}
AFTER = {"repeat": "repeat", "repeat harder": "repeatHarder", "maintenance": "maintenance", "stop": "stop"}


class ParseError(Exception):
    pass


# ----------------------------------------------------------------------------- lengths

LENGTH = r"(\d+(?:\.\d+)?)\s*(km|mi|m|min|mins|minutes|minute|sec|secs|seconds|s|h|hr|hours?)\b"


def length(value, unit):
    """{"meters": …} or {"seconds": …}."""
    v = float(value)
    unit = unit.lower()
    if unit == "km":
        return {"meters": v * 1000}
    if unit == "mi":
        return {"meters": round(v * 1609.344)}
    if unit == "m":
        return {"meters": v}
    if unit in ("min", "mins", "minute", "minutes"):
        return {"seconds": v * 60}
    if unit in ("s", "sec", "secs", "seconds"):
        return {"seconds": v}
    return {"seconds": v * 3600}


def zone_code(token):
    zone = ZONES.get(token.lower())
    if not zone:
        raise ParseError(f"unknown pace '{token}'")
    return zone


# ----------------------------------------------------------------------------- steps

def step(text, recovery=False):
    """One step: '800 m @ I', '1 min E', '2 min @ T effort', '60 s uphill @ hard effort'."""
    text = text.strip().rstrip(",")
    m = re.match(LENGTH, text, re.I)
    if not m:
        raise ParseError(f"no length in '{text}'")
    result = dict(length(m.group(1), m.group(2)))
    rest = text[m.end():].strip()
    uphill = "uphill" in rest.lower() or "hill" in rest.lower()
    effort = bool(re.search(r"\beffort\b", rest, re.I))
    zm = re.search(r"@\s*" + ZONE_PATTERN + r"\b", rest, re.I) or re.match(ZONE_PATTERN + r"\b", rest, re.I)
    if uphill:
        result.update(zone="hill", effort=True, label="Uphill")
    elif zm:
        result.update(zone=zone_code(zm.group(1)), effort=effort)
    elif re.search(r"\b(jog|walk|rest)\b", rest, re.I) or not rest:
        result.update(zone="easy", effort=False)
    else:
        raise ParseError(f"no pace in '{text}'")
    if result["zone"] == "easy" or recovery:
        result["recovery"] = True
    return result


def recovery_step(tail, work):
    """The recovery after each rep: '90 s jog between', 'jog back down (as recovery)'."""
    tail = tail.strip()
    if not tail:
        return None
    if re.search(r"back down", tail, re.I):
        rec = {k: v for k, v in work.items() if k in ("meters", "seconds")}
        rec.update(zone="easy", effort=False, recovery=True, label="Jog back down")
        return rec
    m = re.search(LENGTH, tail, re.I)
    if not m:
        raise ParseError(f"can't read the recovery '{tail}'")
    rec = dict(length(m.group(1), m.group(2)))
    if re.search(r"\bstations?\b", tail, re.I):
        # Station work between reps (Hyrox): no pace, just the clock.
        rec.update(zone="easy", effort=True, recovery=True, label="Station work")
        return rec
    walk = re.search(r"\bwalk\b", tail, re.I)
    rec.update(zone="easy", effort=False, recovery=True, label="Walk" if walk else "Recovery jog")
    return rec


def reps(text):
    """'6 x 800 m @ I, 90 s jog between' or '8 x (1 min @ I / 1 min E)' or '4 x 20 s strides'."""
    m = re.match(r"(\d+)\s*x\s*(.*)$", text.strip(), re.I)
    if not m:
        return None
    count = int(m.group(1))
    body = m.group(2).strip()
    if body.startswith("("):
        close = body.find(")")
        if close < 0:
            raise ParseError("missing ')'")
        inner, tail = body[1:close], body[close + 1:].lstrip(" ,/")
        steps = [step(part, recovery=i > 0 and not re.search("@", part)) for i, part in enumerate(inner.split("/"))]
        # The first step is always work.
        steps[0].pop("recovery", None)
        recovery = recovery_step(tail, steps[0]) if tail else None
        return {"kind": "reps", "repeats": count, "steps": steps, "recovery": recovery}
    if re.search(r"\bstrides?\b", body, re.I):
        lm = re.match(LENGTH, body, re.I)
        work = dict(length(lm.group(1), lm.group(2))) if lm else {"seconds": 20}
        work.update(zone="repetition", effort=True, label="Stride")
        return {"kind": "reps", "repeats": count, "steps": [work],
                "recovery": {"seconds": 60, "zone": "easy", "effort": False, "recovery": True, "label": "Easy jog"}}
    # Work, then the recovery after a comma or a slash.
    parts = re.split(r"\s*[,/]\s*", body, maxsplit=1)
    work = step(parts[0])
    work.pop("recovery", None)
    recovery = recovery_step(parts[1], work) if len(parts) > 1 else None
    return {"kind": "reps", "repeats": count, "steps": [work], "recovery": recovery}


# ----------------------------------------------------------------------------- workouts

def shares(segments):
    """[(weight, zone)] → share parts that add up to 1."""
    total = sum(w for w, _ in segments)
    if total <= 0:
        raise ParseError("no distance to share out")
    return [{"kind": "share", "share": round(w / total, 4), "zone": z} for w, z in segments]


def workout(text, session_type):
    """Parts of a workout. Shares size with the run; reps keep their lengths."""
    raw = (text or "").strip()
    t = raw.replace("×", "x").replace("–", "-").replace("—", "-").replace("->", "→")
    t = re.sub(r"\s+", " ", t)
    plain_zone = "recovery" if session_type == "recovery" else "easy"

    if session_type == "race":
        return [{"kind": "race"}]

    # Progression: '4 km E → 3 km STEADY → 1 km HM'
    if "→" in t:
        segments = []
        for piece in t.split("→"):
            m = re.match(r"\s*" + LENGTH + r"\s*(?:@\s*)?" + ZONE_PATTERN + r"\b", piece, re.I)
            if not m:
                raise ParseError(f"can't read '{piece.strip()}' (expected e.g. '3 km STEADY')")
            amount = length(m.group(1), m.group(2))
            segments.append((amount.get("meters") or amount.get("seconds"), zone_code(m.group(3))))
        return shares(segments)

    # Finish: '13 km E, final 3 km @ STEADY'
    m = re.match(r"(\d+(?:\.\d+)?)\s*km\s*(?:@\s*)?(?:easy|e)\b.*?final\s+(\d+(?:\.\d+)?)\s*km\s*@?\s*" + ZONE_PATTERN + r"\b", t, re.I)
    if m:
        total, final = float(m.group(1)), float(m.group(2))
        if final >= total:
            raise ParseError("the final part is longer than the run")
        return shares([(total - final, plain_zone), (final, zone_code(m.group(3)))])

    # Continuous: '15 min continuous @ T'
    m = re.match(LENGTH + r"\s*(?:continuous\s*)?@\s*" + ZONE_PATTERN + r"\b(\s*effort)?", t, re.I)
    if m and not re.search(r"\d+\s*x\s", t, re.I):
        work = dict(length(m.group(1), m.group(2)))
        work.update(zone=zone_code(m.group(3)), effort=bool(m.group(4)))
        return [{"kind": "reps", "repeats": 1, "steps": [work], "recovery": None}]

    # Easy run with strides: '4 km E + 4 x 20 s strides'
    if "+" in t:
        parts = []
        for piece in [p.strip() for p in t.split("+")]:
            block = reps(piece)
            if block:
                parts.append(block)
            elif re.match(LENGTH + r"\s*(?:@\s*)?(?:easy|e)\b", piece, re.I) or not re.search(r"@", piece):
                parts.append({"kind": "share", "share": 1.0, "zone": plain_zone})
            else:
                raise ParseError(f"can't read '{piece}'")
        return parts

    # Reps: '6 x 800 m @ I, 90 s jog between'
    block = reps(t)
    if block:
        return [block]

    # Anything else without a pace is a plain run: 'E continuous run', 'Very relaxed recovery run'.
    if re.search(r"@", t):
        raise ParseError("has a pace but isn't in a form the app reads")
    return [{"kind": "share", "share": 1.0, "zone": plain_zone}]


def distance(value):
    """Distance column: auto, 25%, 8km."""
    v = "auto" if value is None else str(value).strip().lower()
    if v in ("", "auto"):
        return {"kind": "auto"}
    if v == "race":
        return {"kind": "race"}
    m = re.fullmatch(r"(\d+(?:\.\d+)?)\s*%", v)
    if m:
        return {"kind": "share", "value": float(m.group(1)) / 100}
    if isinstance(value, float) and value <= 1:
        return {"kind": "share", "value": float(value)}
    m = re.fullmatch(r"(\d+(?:\.\d+)?)\s*km", v)
    if m:
        return {"kind": "meters", "value": float(m.group(1)) * 1000}
    raise ParseError(f"distance '{value}' should be auto, a % or km")


# ----------------------------------------------------------------------------- workbook

def rows(ws):
    header = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    for index, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row and row[0] not in (None, ""):
            yield index, dict(zip(header, row))


def levels_from(text):
    """'All levels', 'Beginner only', 'Intermediate and advanced' → the levels named."""
    t = str(text or "all levels").lower()
    if "all" in t:
        return ["beginner", "intermediate", "advanced"]
    return [level for level in ("beginner", "intermediate", "advanced") if level in t]


def number(value, what):
    try:
        return float(value)
    except (TypeError, ValueError):
        raise ParseError(f"{what} should be a number, not '{value}'")


def also_suggested(v, row, fail):
    """Codes of the 'Also suggested' splits (comma-separated names)."""
    text = next((v[k] for k in v if k and str(k).startswith("Also suggested")), None)
    codes = []
    for name in str(text or "").split(","):
        name = name.strip().lower()
        if not name:
            continue
        if name not in SPLITS:
            fail("Programmes", row, f"'Also suggested' names an unknown split: {name}")
            continue
        codes.append(SPLITS[name])
    return codes


def day_focus(name):
    """Push / Pull / Legs / Upper / Lower days; anything else is full body."""
    n = name.lower()
    for word, focus in (("push", "push"), ("pull", "pull"), ("leg", "legs"), ("upper", "upper"), ("lower", "lower"),
                        ("shoulder", "upper"), ("arm", "upper")):
        if word in n:
            return focus
    return "fullBody"


def optional_int(value, what):
    if value in (None, ""):
        return None
    return int(number(value, what))


def strength(wb, fail):
    """Strength splits, their days, and sets/reps by goal."""
    names = {e["name"].lower(): e["name"] for e in json.loads(EXERCISES.read_text())["exercises"]}
    splits = {}
    for row, v in rows(wb["Strength splits"]):
        code = SPLITS.get(str(v["Split"]).strip().lower())
        if not code:
            fail("Strength splits", row, f"unknown split '{v['Split']}'")
            continue
        splits[code] = {"code": code, "name": str(v["Split"]).strip(), "suits": (v.get("Who it suits") or "").strip(),
                        "shown": str(v.get("Shown in onboarding?") or "Yes").strip().lower() != "no", "days": []}
    for row, v in rows(wb["Split days"]):
        code = SPLITS.get(str(v["Split"]).strip().lower())
        if code not in splits:
            if v.get("Day name") or v.get("Exercise"):
                fail("Split days", row, f"unknown split '{v['Split']}'")
            continue  # a note under the table
        day_name = str(v.get("Day name") or "").strip()
        exercise = str(v.get("Exercise") or "").strip()
        if not day_name or not exercise:
            fail("Split days", row, "needs a day name and an exercise")
            continue
        try:
            slot = {"role": "main" if str(v.get("Role") or "").strip().lower() == "main" else "accessory"}
            if exercise.lower() in PATTERNS:
                slot["pattern"] = PATTERNS[exercise.lower()]
            elif exercise.lower() in names:
                slot["pattern"] = "named"
                slot["exerciseName"] = names[exercise.lower()]
            else:
                raise ParseError(f"unknown exercise '{exercise}' (pick one from the drop-down)")
            for key, column in (("sets", "Sets"), ("repsLow", "Reps (low)"), ("repsHigh", "Reps (high)"), ("restSeconds", "Rest (s)")):
                value = optional_int(v.get(column), column)
                if value is not None:
                    slot[key] = value
            if v.get("Notes"):
                slot["note"] = str(v["Notes"]).strip()
            order = optional_int(v.get("Order"), "Order") or 0
        except ParseError as e:
            fail("Split days", row, str(e))
            continue
        days = splits[code]["days"]
        day = next((d for d in days if d["name"] == day_name), None)
        if day is None:
            day = {"name": day_name, "focus": day_focus(day_name), "slots": []}
            days.append(day)
        day["slots"].append((order, slot))
    for split in splits.values():
        for day in split["days"]:
            day["slots"] = [slot for _, slot in sorted(day["slots"], key=lambda item: item[0])]
    by_goal = {}
    for row, v in rows(wb["Strength by goal"]):
        goal = GOALS.get(str(v["Programme"]).strip().lower())
        role = str(v.get("Role") or "").strip().lower()
        if not goal or role not in ("main", "accessory"):
            continue
        try:
            by_goal.setdefault(goal, {})[role] = {
                "sets": int(number(v.get("Sets"), "Sets")), "repsLow": int(number(v.get("Reps (low)"), "Reps (low)")),
                "repsHigh": int(number(v.get("Reps (high)"), "Reps (high)")), "restSeconds": int(number(v.get("Rest (s)"), "Rest (s)")),
            }
        except ParseError as e:
            fail("Strength by goal", row, str(e))
    by_goal = {goal: roles for goal, roles in by_goal.items() if "main" in roles and "accessory" in roles}
    return {"splits": [s for s in splits.values() if s["days"]], "byGoal": by_goal}


def convert(path):
    wb = load_workbook(path, data_only=True)
    errors = []

    def fail(tab, row, message):
        errors.append(f"{tab} row {row}: {message}")

    # Rules
    rules = {"beginner": {}, "intermediate": {}, "advanced": {}}
    rule_keys = [
        ("starting volume if they don't run", "defaultStartMeters", 1000), ("minimum starting volume", "minStartMeters", 1000),
        ("maximum starting volume", "maxStartMeters", 1000), ("longest single run", "longestRunMeters", 1000),
        ("shortest run", "shortestRunMeters", 1000), ("max weekly increase", "maxWeeklyIncrease", 1),
        ("long run share", "longRunShare", 1),
    ]
    # Sizing rules (optional: the app has the same defaults).
    optional_keys = [
        ("warm-up and cool-down", "warmUpShare", 1), ("warm-up minimum", "warmUpMinMeters", 1000),
        ("warm-up maximum", "warmUpMaxMeters", 1000), ("fast reps max", "fastRepsShare", 1),
        ("fast reps floor", "fastRepsFloorMeters", 1000), ("threshold reps max", "thresholdRepsShare", 1),
        ("threshold reps floor", "thresholdRepsFloorMeters", 1000), ("steady reps max", "steadyRepsShare", 1),
        ("steady reps floor", "steadyRepsFloorMeters", 1000), ("recovery run max", "recoveryRunOfLong", 1),
        ("easy run max", "easyRunOfLong", 1), ("shakeout max", "shakeoutOfWeek", 1),
        ("recovery week after", "recoveryWeekShare", 1), ("ease down per week", "easeDownPerWeek", 1),
        ("easy first week after a race", "easyStartShare", 1),
    ]
    for row, values in rows(wb["Rules"]):
        name = str(values.get("Rule", "")).lower()
        for prefix, key, scale in rule_keys + optional_keys:
            if name.startswith(prefix):
                for level in rules:
                    try:
                        rules[level][key] = number(values.get(level.capitalize()), f"Rules {level}") * scale
                    except ParseError as e:
                        fail("Rules", row, str(e))
    for level, values in rules.items():
        missing = [key for _, key, _ in rule_keys if key not in values]
        if missing:
            fail("Rules", "-", f"missing {', '.join(missing)} for {level}")

    # Programmes (only those with weeks written) and which runners see each goal.
    programmes = {}
    availability = {}
    goal_settings = {}
    for row, v in rows(wb["Programmes"]):
        goal = GOALS.get(str(v["Programme"]).strip().lower())
        if not goal:
            continue
        levels = levels_from(v.get("Available to"))
        if not levels:
            fail("Programmes", row, f"'Available to' should name beginner, intermediate and/or advanced, or say all levels")
        availability[goal] = levels
        try:
            goal_settings[goal] = {
                "runsMin": int(number(v.get("Runs per week (min)"), "Runs per week (min)")),
                "runsMax": int(number(v.get("Runs per week (max)"), "Runs per week (max)")),
                "liftsMin": int(number(v.get("Lifts per week (min)"), "Lifts per week (min)")),
                "liftsMax": int(number(v.get("Lifts per week (max)"), "Lifts per week (max)")),
                "defaultSplit": SPLITS.get(str(v.get("Default split") or "").strip().lower()),
                "alsoSuggested": also_suggested(v, row, fail),
            }
        except ParseError as e:
            fail("Programmes", row, str(e))
        try:
            programmes[goal] = {
                "goal": goal,
                "name": str(v.get("Name shown in app") or v["Programme"]).strip(),
                "description": (v.get("Description shown in onboarding") or "").strip(),
                "lengthWeeks": int(number(v.get("Length (weeks)"), "Length")),
                "afterLastWeek": AFTER.get(str(v.get("After the last week") or "repeat").strip().lower(), "repeat"),
                "runsPerWeekMin": int(number(v.get("Runs per week (min)"), "Runs per week (min)")),
                "runsPerWeekMax": int(number(v.get("Runs per week (max)"), "Runs per week (max)")),
                "longRunCapMeters": number(v.get("Long-run cap (km)"), "Long-run cap") * 1000,
                "weeklyCapMeters": number(v.get("Weekly volume cap (km)"), "Weekly volume cap") * 1000
                    if v.get("Weekly volume cap (km)") not in (None, "") else None,
                "taperWeeks": int(number(v.get("Taper weeks") or 0, "Taper weeks")),
                "raceReadyMeters": {
                    level: number(v.get(f"Race-ready volume: {level} (km/week)"), f"Race-ready volume ({level})") * 1000
                    for level in ("beginner", "intermediate", "advanced")
                    if v.get(f"Race-ready volume: {level} (km/week)") not in (None, "")
                },
                "weeks": [],
                "sessions": [],
            }
        except ParseError as e:
            fail("Programmes", row, str(e))

    # Run weeks
    last_phase = {}
    for row, v in rows(wb["Run weeks"]):
        goal = GOALS.get(str(v["Programme"]).strip().lower())
        if goal not in programmes:
            fail("Run weeks", row, f"unknown programme '{v['Programme']}'")
            continue
        try:
            phase_text = str(v.get("Phase") or "").strip().lower()
            if phase_text not in PHASES:
                raise ParseError(f"unknown phase '{v.get('Phase')}'")
            deload = phase_text == "deload"
            phase = PHASES[phase_text] or last_phase.get(goal, "base")
            last_phase[goal] = phase
            long_run = distance(v.get("Long run"))
            if long_run["kind"] == "share" and long_run["value"] > 1:
                raise ParseError("long run % should be below 100%")
            emphasis = EMPHASIS.get(str(v.get("Strength emphasis") or "normal").strip().lower())
            if emphasis is None:
                raise ParseError(f"unknown strength emphasis '{v.get('Strength emphasis')}'")
            volume = v.get("Volume (% of race-ready)", v.get("Volume (% of starting)"))
            volume = number(volume.rstrip("%"), "Volume") / 100 if isinstance(volume, str) else number(volume, "Volume")
            note = v.get("Notes for the runner")
            programmes[goal]["weeks"].append({
                "week": int(number(v.get("Week"), "Week")), "phase": phase, "isDeload": deload,
                "isRaceWeek": phase_text == "race week", "volume": round(volume, 4), "longRun": long_run,
                "strength": emphasis, "note": str(note).strip() if note else None,
            })
        except ParseError as e:
            fail("Run weeks", row, str(e))

    # Run sessions
    for row, v in rows(wb["Run sessions"]):
        goal = GOALS.get(str(v["Programme"]).strip().lower())
        if goal not in programmes:
            fail("Run sessions", row, f"unknown programme '{v['Programme']}'")
            continue
        try:
            type_text = str(v.get("Session type") or "").strip().lower()
            if type_text not in SESSION_TYPES:
                raise ParseError(f"unknown session type '{v.get('Session type')}'")
            session_type, name = SESSION_TYPES[type_text]
            text = v.get("Workout")
            note = v.get("Notes for the runner")
            programmes[goal]["sessions"].append({
                "week": int(number(v.get("Week"), "Week")),
                "order": int(number(v.get("Order in week") or 0, "Order in week")),
                "name": name, "type": session_type,
                "distance": distance(v.get("Distance")),
                "parts": workout(text, session_type),
                "text": str(text).strip() if text else None,
                "note": str(note).strip() if note else None,
            })
        except ParseError as e:
            fail("Run sessions", row, f"{e}  [{v.get('Workout')}]")

    # Checks across tabs
    written = []
    for goal, p in programmes.items():
        if not p["weeks"]:
            continue
        weeks = sorted(w["week"] for w in p["weeks"])
        if weeks != list(range(1, len(weeks) + 1)):
            fail("Run weeks", "-", f"{p['name']}: weeks should run 1, 2, 3… without gaps ({weeks})")
        if len(weeks) != p["lengthWeeks"]:
            fail("Programmes", "-", f"{p['name']}: length is {p['lengthWeeks']} weeks but {len(weeks)} are written")
        for week in weeks:
            if not any(s["week"] == week for s in p["sessions"]):
                fail("Run sessions", "-", f"{p['name']} week {week} has no sessions")
            if sum(1 for s in p["sessions"] if s["week"] == week and s["type"] == "long") > 1:
                fail("Run sessions", "-", f"{p['name']} week {week} has more than one long run")
        p["weeks"].sort(key=lambda w: w["week"])
        p["sessions"].sort(key=lambda s: (s["week"], s["order"]))
        written.append(p)

    return {"version": 1, "rules": rules, "availability": availability, "goals": goal_settings,
            "strength": strength(wb, fail), "programmes": written}, errors


def main():
    path = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_WORKBOOK
    data, errors = convert(path)
    if errors:
        print("Not written: fix these in the workbook first.\n")
        for e in errors:
            print("  - " + e)
        sys.exit(1)
    OUTPUT.write_text(json.dumps(data, indent=1, ensure_ascii=False) + "\n")
    for p in data["programmes"]:
        print(f"{p['name']}: {len(p['weeks'])} weeks, {len(p['sessions'])} sessions")
    print(f"Wrote {OUTPUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
